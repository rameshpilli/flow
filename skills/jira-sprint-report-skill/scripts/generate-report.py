#!/usr/bin/env python3
"""
JIRA Sprint Report Generator
Generates a PowerPoint presentation from sprint data JSON

Usage:
    python generate-report.py <input.json> [output.pptx] [--template <template.pptx>]

Template Support:
    - Pass --template to use an existing branded PowerPoint as the base
    - The template's master slides, backgrounds, and branding will be preserved
    - New slides are added using the template's slide layouts
"""

import json
import sys
import os
import argparse
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData


# Color palette (RGB tuples)
class Colors:
    PRIMARY = RGBColor(0x1a, 0x5f, 0x7a)
    PRIMARY_DARK = RGBColor(0x0d, 0x3d, 0x4d)
    SUCCESS = RGBColor(0x22, 0xc5, 0x5e)
    WARNING = RGBColor(0xf5, 0x9e, 0x0b)
    DANGER = RGBColor(0xef, 0x44, 0x44)
    INFO = RGBColor(0x3b, 0x82, 0xf6)
    PURPLE = RGBColor(0x8b, 0x5c, 0xf6)
    SURFACE = RGBColor(0xf1, 0xf5, 0xf9)
    WHITE = RGBColor(0xff, 0xff, 0xff)
    DARK = RGBColor(0x1e, 0x29, 0x3b)
    MUTED = RGBColor(0x64, 0x74, 0x8b)
    MUTED_LIGHT = RGBColor(0x94, 0xa3, 0xb8)
    STORY = RGBColor(0x63, 0x66, 0xf1)
    BUG = RGBColor(0xef, 0x44, 0x44)
    ROW_ALT = RGBColor(0xf8, 0xfa, 0xfc)


# Slide dimensions (16:9)
SLIDE_WIDTH = Inches(10)
SLIDE_HEIGHT = Inches(5.625)

# Table settings
ISSUES_PER_PAGE = 20
COL_WIDTHS = [Inches(0.9), Inches(0.7), Inches(4.0), Inches(0.5), Inches(1.7), Inches(0.9), Inches(0.9)]


def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None):
    """Add a shape to a slide with optional fill and line colors."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=12, font_bold=False,
                 font_color=Colors.DARK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
                 font_name="Arial"):
    """Add a text box to a slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(font_size)
    p.font.bold = font_bold
    p.font.color.rgb = font_color
    p.font.name = font_name
    p.alignment = align
    tf.anchor = valign
    return txBox


def create_title_slide(prs, data, use_template=False):
    """Create the title slide."""
    # Use blank layout or first layout from template
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    if not use_template:
        # Add background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT
        )
        background.fill.solid()
        background.fill.fore_color.rgb = Colors.PRIMARY
        background.line.fill.background()

    # Title
    title_color = Colors.DARK if use_template else Colors.WHITE
    add_text_box(slide, Inches(0.5), Inches(1.8), Inches(9), Inches(0.8),
                 data['sprint']['name'], font_size=48, font_bold=True,
                 font_color=title_color, align=PP_ALIGN.CENTER)

    # Subtitle
    subtitle_color = Colors.MUTED if use_template else Colors.MUTED_LIGHT
    add_text_box(slide, Inches(0.5), Inches(2.6), Inches(9), Inches(0.5),
                 "Summary Report", font_size=20, font_color=subtitle_color,
                 align=PP_ALIGN.CENTER)

    # Info cards
    cards = [
        {"label": "Release Start", "value": data['sprint']['startDate']},
        {"label": "Release End", "value": data['sprint']['endDate']},
        {"label": "Team Size", "value": f"{data['sprint']['teamSize']} Members"}
    ]

    card_w, card_h, card_gap, card_y = Inches(2), Inches(0.8), Inches(0.3), Inches(3.4)
    start_x = (SLIDE_WIDTH - (3 * card_w + 2 * card_gap)) / 2

    for i, card in enumerate(cards):
        x = start_x + i * (card_w + card_gap)

        # Card background
        card_shape = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_w, card_h,
                              fill_color=Colors.WHITE)
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = Colors.WHITE

        # Label
        add_text_box(slide, x, card_y + Inches(0.08), card_w, Inches(0.25),
                    card['label'], font_size=9, font_color=Colors.MUTED, align=PP_ALIGN.CENTER)

        # Value
        value_color = Colors.DARK if use_template else Colors.PRIMARY_DARK
        add_text_box(slide, x, card_y + Inches(0.35), card_w, Inches(0.35),
                    card['value'], font_size=14, font_bold=True, font_color=value_color,
                    align=PP_ALIGN.CENTER)

    # Team name
    add_text_box(slide, Inches(0.5), Inches(4.5), Inches(9), Inches(0.4),
                 data['sprint'].get('teamName', 'Engineering Team'), font_size=12,
                 font_color=subtitle_color, align=PP_ALIGN.CENTER)


def create_dashboard_slide(prs, data, use_template=False):
    """Create the dashboard slide with KPIs and charts."""
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    if not use_template:
        # Background
        bg = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT,
                      fill_color=Colors.SURFACE)

    # Header bar
    header = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.55),
                      fill_color=Colors.PRIMARY)

    add_text_box(slide, Inches(0.3), Inches(0.12), Inches(5), Inches(0.35),
                 f"{data['sprint']['name']} Dashboard", font_size=14, font_bold=True,
                 font_color=Colors.WHITE)

    # Completion badge
    total_issues = data['metrics']['completed']['issues'] + data['metrics']['spillover']['issues']
    completion_rate = round((data['metrics']['completed']['issues'] / total_issues) * 100)

    badge = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8), Inches(0.1),
                     Inches(1.8), Inches(0.35), fill_color=Colors.WHITE)
    add_text_box(slide, Inches(8), Inches(0.12), Inches(1.8), Inches(0.32),
                 f"{completion_rate}% Complete", font_size=10, font_bold=True,
                 font_color=Colors.SUCCESS, align=PP_ALIGN.CENTER)

    # KPI Cards
    kpis = [
        {"label": "COMPLETED", "value": data['metrics']['completed']['points'],
         "unit": "pts", "sub": f"{data['metrics']['completed']['issues']} issues", "color": Colors.SUCCESS},
        {"label": "BLOCKED", "value": data['metrics']['blocked']['issues'],
         "unit": "issues", "sub": "Needs attention", "color": Colors.DANGER, "sub_color": Colors.DANGER},
        {"label": "SPILLOVER", "value": data['metrics']['spillover']['points'],
         "unit": "pts", "sub": f"{data['metrics']['spillover']['issues']} issues", "color": Colors.WARNING},
        {"label": "AVG VELOCITY", "value": data['metrics']['avgVelocity'],
         "unit": "pts", "sub": "Last 5 sprints", "color": Colors.PURPLE}
    ]

    kpi_y, kpi_h, kpi_w, kpi_gap = Inches(0.7), Inches(1.1), Inches(2.3), Inches(0.15)
    kpi_start_x = Inches(0.2)

    for i, kpi in enumerate(kpis):
        x = kpi_start_x + i * (kpi_w + kpi_gap)

        # Card background
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, kpi_y, kpi_w, kpi_h, fill_color=Colors.WHITE)

        # Colored accent bar
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, kpi_y, Inches(0.06), kpi_h, fill_color=kpi['color'])

        # Label
        add_text_box(slide, x + Inches(0.15), kpi_y + Inches(0.1), kpi_w - Inches(0.25), Inches(0.22),
                    kpi['label'], font_size=9, font_color=Colors.MUTED)

        # Value
        add_text_box(slide, x + Inches(0.15), kpi_y + Inches(0.35), Inches(1.4), Inches(0.5),
                    str(kpi['value']), font_size=28, font_bold=True, font_color=Colors.DARK)

        # Unit
        add_text_box(slide, x + Inches(1.4), kpi_y + Inches(0.52), Inches(0.8), Inches(0.3),
                    kpi['unit'], font_size=11, font_color=Colors.MUTED)

        # Sub text
        sub_color = kpi.get('sub_color', Colors.MUTED_LIGHT)
        add_text_box(slide, x + Inches(0.15), kpi_y + Inches(0.85), kpi_w - Inches(0.25), Inches(0.2),
                    kpi['sub'], font_size=9, font_color=sub_color)

    # Charts section
    chart_y, chart_h = Inches(2.0), Inches(3.0)

    # Velocity chart background
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.2), chart_y, Inches(5.8), chart_h,
             fill_color=Colors.WHITE)
    add_text_box(slide, Inches(0.35), chart_y + Inches(0.12), Inches(5.5), Inches(0.28),
                "Velocity Comparison: Current vs Previous Sprint", font_size=11,
                font_bold=True, font_color=Colors.DARK)

    # Create velocity line chart
    chart_data = CategoryChartData()
    chart_data.categories = data['velocity']['labels']
    current_total = data['velocity']['current'][-1]
    prev_total = data['velocity']['previous'][-1]
    chart_data.add_series(f'Current Sprint ({current_total} pts)', data['velocity']['current'])
    chart_data.add_series(f'Previous Sprint ({prev_total} pts)', data['velocity']['previous'])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(0.3), chart_y + Inches(0.5),
        Inches(5.6), Inches(2.3), chart_data
    ).chart

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    # Status pie chart background
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(6.15), chart_y, Inches(3.65), chart_h,
             fill_color=Colors.WHITE)
    add_text_box(slide, Inches(6.3), chart_y + Inches(0.12), Inches(3.35), Inches(0.28),
                "Status Distribution", font_size=11, font_bold=True, font_color=Colors.DARK)

    # Count statuses
    status_counts = {"Done": 0, "In Progress": 0, "Blocked": 0}
    for issue in data['issues']:
        if issue['status'] in status_counts:
            status_counts[issue['status']] += 1

    # Create pie chart
    pie_data = CategoryChartData()
    pie_data.categories = [f"Done ({status_counts['Done']})",
                          f"In Progress ({status_counts['In Progress']})",
                          f"Blocked ({status_counts['Blocked']})"]
    pie_data.add_series('Status', [status_counts['Done'], status_counts['In Progress'], status_counts['Blocked']])

    pie_chart = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(6.2), chart_y + Inches(0.5),
        Inches(3.55), Inches(2.3), pie_data
    ).chart

    pie_chart.has_legend = True
    pie_chart.legend.position = XL_LEGEND_POSITION.BOTTOM


def create_issue_slide(prs, issues, page_num, total_pages, story_count, bug_count, use_template=False):
    """Create an issue table slide."""
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    if not use_template:
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT,
                 fill_color=Colors.SURFACE)

    # Header
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.5),
             fill_color=Colors.PRIMARY)
    add_text_box(slide, Inches(0.3), Inches(0.1), Inches(5), Inches(0.32),
                f"Sprint Issues: Stories & Bugs ({page_num}/{total_pages})",
                font_size=13, font_bold=True, font_color=Colors.WHITE)

    # Badges
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(0.08),
             Inches(1.3), Inches(0.32), fill_color=Colors.STORY)
    add_text_box(slide, Inches(7.0), Inches(0.1), Inches(1.3), Inches(0.3),
                f"{story_count} Stories", font_size=9, font_bold=True,
                font_color=Colors.WHITE, align=PP_ALIGN.CENTER)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(0.08),
             Inches(1.3), Inches(0.32), fill_color=Colors.BUG)
    add_text_box(slide, Inches(8.45), Inches(0.1), Inches(1.3), Inches(0.3),
                f"{bug_count} Bugs", font_size=9, font_bold=True,
                font_color=Colors.WHITE, align=PP_ALIGN.CENTER)

    # Table
    table_y = Inches(0.6)
    row_h = Inches(0.23)

    rows = len(issues) + 1  # +1 for header
    cols = 7

    table = slide.shapes.add_table(rows, cols, Inches(0.2), table_y,
                                   sum(COL_WIDTHS), row_h * rows).table

    # Set column widths
    for i, width in enumerate(COL_WIDTHS):
        table.columns[i].width = width

    # Header row
    headers = ["Ticket", "Type", "Summary", "Est", "Assignee", "Status", "Days +/-"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = Colors.PRIMARY_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = Colors.WHITE
        p.alignment = PP_ALIGN.CENTER if i != 2 else PP_ALIGN.LEFT

    # Data rows
    for row_idx, issue in enumerate(issues):
        row_fill = Colors.ROW_ALT if row_idx % 2 == 0 else Colors.WHITE
        type_color = Colors.STORY if issue['type'] == "Story" else Colors.BUG

        status_color = Colors.SUCCESS
        if issue['status'] == "In Progress":
            status_color = Colors.WARNING
        elif issue['status'] == "Blocked":
            status_color = Colors.DANGER

        days_over = issue.get('daysOver')
        if days_over is not None:
            if days_over < 0:
                days_text, days_color = f"{days_over}d", Colors.SUCCESS
            elif days_over == 0:
                days_text, days_color = "0d", Colors.INFO
            else:
                days_text, days_color = f"+{days_over}d", Colors.DANGER
        else:
            days_text, days_color = "-", Colors.MUTED

        row_data = [
            (issue['key'], Colors.DARK, PP_ALIGN.CENTER),
            (issue['type'], type_color, PP_ALIGN.CENTER),
            (issue['summary'], Colors.DARK, PP_ALIGN.LEFT),
            (str(issue['estimate']), Colors.DARK, PP_ALIGN.CENTER),
            (issue['assignee'], Colors.MUTED, PP_ALIGN.LEFT),
            (issue['status'], status_color, PP_ALIGN.CENTER),
            (days_text, days_color, PP_ALIGN.CENTER)
        ]

        for col_idx, (text, color, align) in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_fill
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(8)
            p.font.color.rgb = color
            p.font.bold = col_idx in [1, 5, 6]  # Type, Status, Days
            p.alignment = align

    # Legend on first page
    if page_num == 1:
        legend_y = table_y + row_h * rows + Inches(0.15)
        add_text_box(slide, Inches(0.3), legend_y, Inches(0.6), Inches(0.2),
                    "Days +/-:", font_size=8, font_bold=True, font_color=Colors.DARK)
        add_text_box(slide, Inches(0.95), legend_y, Inches(1.3), Inches(0.2),
                    "-2d = 2 days early", font_size=8, font_color=Colors.SUCCESS)
        add_text_box(slide, Inches(2.35), legend_y, Inches(1), Inches(0.2),
                    "0d = on time", font_size=8, font_color=Colors.INFO)
        add_text_box(slide, Inches(3.45), legend_y, Inches(1.4), Inches(0.2),
                    "+3d = 3 days over", font_size=8, font_color=Colors.DANGER)


def create_summary_slide(prs, data, use_template=False):
    """Create the summary slide."""
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    if not use_template:
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT,
                 fill_color=Colors.SURFACE)

    # Header
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.5),
             fill_color=Colors.PRIMARY)
    add_text_box(slide, Inches(0.3), Inches(0.1), Inches(6), Inches(0.32),
                "Summary & Next Sprint", font_size=13, font_bold=True, font_color=Colors.WHITE)

    left_w = Inches(5.8)

    # Key Achievements
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.6), left_w, Inches(2.4),
             fill_color=Colors.WHITE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.6), Inches(0.06), Inches(2.4),
             fill_color=Colors.SUCCESS)
    add_text_box(slide, Inches(0.35), Inches(0.68), left_w - Inches(0.3), Inches(0.25),
                "✅ Key Achievements", font_size=12, font_bold=True, font_color=Colors.DARK)

    achievements = data.get('achievements', [])[:5]
    for i, achievement in enumerate(achievements):
        add_text_box(slide, Inches(0.4), Inches(1.0) + i * Inches(0.36),
                    left_w - Inches(0.4), Inches(0.34),
                    f"•  {achievement}", font_size=10, font_color=Colors.DARK,
                    valign=MSO_ANCHOR.TOP)

    # Blockers
    blockers = [i for i in data['issues'] if i['status'] == 'Blocked']
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(3.15), left_w, Inches(1.85),
             fill_color=Colors.WHITE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(3.15), Inches(0.06), Inches(1.85),
             fill_color=Colors.DANGER)
    add_text_box(slide, Inches(0.35), Inches(3.23), left_w - Inches(0.3), Inches(0.25),
                "🚫 Current Blockers", font_size=12, font_bold=True, font_color=Colors.DARK)

    for i, issue in enumerate(blockers[:3]):
        add_text_box(slide, Inches(0.4), Inches(3.58) + i * Inches(0.45),
                    left_w - Inches(0.5), Inches(0.2),
                    f"{issue['key']}: {issue['summary']}", font_size=10, font_bold=True,
                    font_color=Colors.DARK)
        add_text_box(slide, Inches(0.4), Inches(3.8) + i * Inches(0.45),
                    left_w - Inches(0.5), Inches(0.18),
                    f"Blocker: {issue.get('blocker', 'Not specified')}", font_size=9,
                    font_color=Colors.DANGER)

    # Next Sprint card
    right_x, right_w = Inches(6.15), Inches(3.65)
    ns = data.get('nextSprint', {})

    add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(0.6), right_w, Inches(4.4),
             fill_color=Colors.PRIMARY)

    add_text_box(slide, right_x + Inches(0.2), Inches(0.75), right_w - Inches(0.4), Inches(0.3),
                "📅 Next Sprint", font_size=14, font_bold=True, font_color=Colors.WHITE)
    add_text_box(slide, right_x + Inches(0.2), Inches(1.1), right_w - Inches(0.4), Inches(0.4),
                ns.get('name', 'TBD'), font_size=28, font_bold=True, font_color=Colors.WHITE)

    # Divider
    add_shape(slide, MSO_SHAPE.RECTANGLE, right_x + Inches(0.2), Inches(1.6),
             right_w - Inches(0.4), Inches(0.02), fill_color=Colors.MUTED_LIGHT)

    info = [
        ("Dates", ns.get('dates', 'TBD'), Colors.WHITE),
        ("Duration", ns.get('duration', 'TBD'), Colors.WHITE),
        ("Team Size", f"{data['sprint']['teamSize']} members", Colors.WHITE),
        ("Planned Capacity", f"{ns.get('plannedCapacity', 0)} story points", Colors.WHITE),
        ("Spillover", f"{ns.get('spillover', 0)} pts", Colors.WARNING),
        ("Available for New Work", f"{ns.get('available', 0)} pts", Colors.SUCCESS)
    ]

    for i, (label, value, color) in enumerate(info):
        add_text_box(slide, right_x + Inches(0.2), Inches(1.8) + i * Inches(0.38),
                    right_w - Inches(0.4), Inches(0.2),
                    label, font_size=9, font_color=Colors.MUTED_LIGHT)
        add_text_box(slide, right_x + Inches(0.2), Inches(2.0) + i * Inches(0.38),
                    right_w - Inches(0.4), Inches(0.22),
                    value, font_size=11, font_bold=True, font_color=color)

    # Sprint goal
    add_shape(slide, MSO_SHAPE.RECTANGLE, right_x + Inches(0.2), Inches(4.15),
             right_w - Inches(0.4), Inches(0.02), fill_color=Colors.MUTED_LIGHT)
    add_text_box(slide, right_x + Inches(0.2), Inches(4.3), right_w - Inches(0.4), Inches(0.2),
                "Sprint Goal", font_size=9, font_color=Colors.MUTED_LIGHT)
    add_text_box(slide, right_x + Inches(0.2), Inches(4.55), right_w - Inches(0.4), Inches(0.6),
                ns.get('goal', 'TBD'), font_size=10, font_color=Colors.WHITE,
                valign=MSO_ANCHOR.TOP)


def generate_report(data, output_path, template_path=None):
    """Generate the PowerPoint report."""
    use_template = template_path is not None and os.path.exists(template_path)

    if use_template:
        prs = Presentation(template_path)
        print(f"  Using template: {template_path}")
    else:
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT

    # Slide 1: Title
    create_title_slide(prs, data, use_template)

    # Slide 2: Dashboard
    create_dashboard_slide(prs, data, use_template)

    # Slides 3-N: Issues
    sorted_issues = sorted(data['issues'], key=lambda x: (
        {"Done": 0, "In Progress": 1, "Blocked": 2}.get(x['status'], 3),
        {"Story": 0, "Bug": 1}.get(x['type'], 2)
    ))

    story_count = len([i for i in data['issues'] if i['type'] == 'Story'])
    bug_count = len([i for i in data['issues'] if i['type'] == 'Bug'])
    total_pages = (len(sorted_issues) + ISSUES_PER_PAGE - 1) // ISSUES_PER_PAGE

    for page in range(total_pages):
        start = page * ISSUES_PER_PAGE
        page_issues = sorted_issues[start:start + ISSUES_PER_PAGE]
        create_issue_slide(prs, page_issues, page + 1, total_pages, story_count, bug_count, use_template)

    # Last slide: Summary
    create_summary_slide(prs, data, use_template)

    # Save
    prs.save(output_path)
    return 2 + total_pages + 1  # Title + Dashboard + Issue pages + Summary


def main():
    parser = argparse.ArgumentParser(
        description='JIRA Sprint Report Generator',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
Examples:
  python generate-report.py sprint-data.json
  python generate-report.py sprint-data.json Report.pptx
  python generate-report.py sprint-data.json Report.pptx --template BASELINE_TEMPLATE.pptx

Template Support:
  Use --template to apply your branded PowerPoint template.
  The template's master slides, backgrounds, and branding will be preserved.
        '''
    )
    parser.add_argument('input', help='Input JSON file with sprint data')
    parser.add_argument('output', nargs='?', help='Output PPTX file (optional)')
    parser.add_argument('--template', '-t', help='Branded PPTX template file')

    args = parser.parse_args()

    input_path = args.input
    output_path = args.output or input_path.replace('.json', '_Report.pptx')
    template_path = args.template

    if not os.path.exists(input_path):
        print(f"Error: Input file not found: {input_path}")
        sys.exit(1)

    if template_path and not os.path.exists(template_path):
        print(f"Warning: Template file not found: {template_path}")
        print("Generating without template...")
        template_path = None

    print(f"Reading data from: {input_path}")
    with open(input_path, 'r') as f:
        data = json.load(f)

    print("Generating report...")
    print(f"  Sprint: {data['sprint']['name']}")
    print(f"  Team: {data['sprint'].get('teamName', 'Unknown')} ({data['sprint']['teamSize']} members)")
    story_count = len([i for i in data['issues'] if i['type'] == 'Story'])
    bug_count = len([i for i in data['issues'] if i['type'] == 'Bug'])
    print(f"  Issues: {len(data['issues'])} ({story_count} Stories, {bug_count} Bugs)")

    total_slides = generate_report(data, output_path, template_path)

    print(f"\n✅ Report saved: {output_path}")
    print(f"   Total slides: {total_slides}")


if __name__ == '__main__':
    main()
